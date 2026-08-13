"""
Script de Consolidacion y Calculo de Rendimientos para Buses
Autor: Pablo Reyes
Descripcion: Procesa multiples archivos Excel consolidados diarios, los integra con 
             la plantilla de rangos de flota y genera un reporte final con formulas 
             activas y formato condicional en Excel.
"""

import os
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.formatting.rule import FormulaRule

# --- CONFIGURACION DE RUTAS ---
DIR_RANGOS = "entrada_rangos"
DIR_CONSOLIDADOS = "entrada_consolidados"
DIR_SALIDA = "Salida"
ARCHIVO_SALIDA = os.path.join(DIR_SALIDA, "REPORTE_RENDIMIENTO.xlsx")


def crear_carpetas():
    """Crea los directorios de trabajo si no existen."""
    os.makedirs(DIR_RANGOS, exist_ok=True)
    os.makedirs(DIR_CONSOLIDADOS, exist_ok=True)
    os.makedirs(DIR_SALIDA, exist_ok=True)


def limpiar_hora(valor):
    """Estandariza el formato de hora a HH:MM."""
    if pd.isna(valor):
        return ''
    valor = str(valor).strip().upper()
    valor = re.sub(r'[.,;]', ':', valor)

    if re.match(r'^\d{1,2}:\d{2}:\d{2}$', valor):  # Formato HH:MM:SS
        return valor[:5]
    if re.match(r'^\d{1,2}:\d{2}$', valor):        # Formato HH:MM
        return valor.zfill(5)
    if valor.isdigit() and len(valor) in [3, 4]:   # Ej: 830 -> 08:30
        valor = valor.zfill(4)
        return f"{valor[:2]}:{valor[2:]}"
    
    return valor


def cargar_rangos():
    """Carga y une las hojas 'detalle' y 'rango' del archivo Rangos.xlsx."""
    detalle_path = os.path.join(DIR_RANGOS, "Rangos.xlsx")
    
    if not os.path.exists(detalle_path):
        print(f"[ERROR] No se encontro el archivo de rangos en: {detalle_path}")
        return None

    try:
        df_detalle = pd.read_excel(detalle_path, sheet_name="detalle")
        df_rango = pd.read_excel(detalle_path, sheet_name="rango")

        # Normalizar nombres de columnas
        df_detalle.columns = df_detalle.columns.str.strip().str.upper()
        df_rango.columns = df_rango.columns.str.strip().str.upper()

        # Limpieza de claves principales
        df_detalle['PATENTE'] = df_detalle['PATENTE'].astype(str).str.strip().str.upper()
        df_detalle['N INTERNO'] = df_detalle['N INTERNO'].astype(str).str.strip()

        # Renombrar columnas de rango
        df_rango.rename(columns={
            "RANGO_MIN": "DESDE",
            "RANGO_MAX": "HASTA"
        }, inplace=True)

        if "NORMA" not in df_rango.columns:
            raise ValueError("La hoja 'rango' debe contener la columna 'NORMA'.")

        # Merge por MODELO y NORMA
        df_merge = pd.merge(
            df_detalle,
            df_rango[['MODELO', 'NORMA', 'DESDE', 'HASTA']],
            on=['MODELO', 'NORMA'],
            how='left'
        )

        rangos_df = df_merge.drop_duplicates()
        print("[OK] Rangos integrados correctamente desde 'detalle' + 'rango'")
        return rangos_df

    except Exception as e:
        print(f"[ERROR] Ocurrio un problema cargando Rangos.xlsx: {e}")
        return None


def cargar_consolidados():
    """Lee y concatena todos los archivos Excel validos en la carpeta de consolidados."""
    archivos = [f for f in os.listdir(DIR_CONSOLIDADOS) if f.endswith(('.xlsx', '.xlsm'))]
    dataframes = []

    columnas_requeridas = {
        'NUMERO INTERNO', 'PATENTE', 'ODOMETRO', 'LITROS',
        'TERMINAL', 'FOLIO', 'FECHA PLANILLA', 'FECHA REAL',
        'ROL', 'HORA', 'BOMBERO', 'TURNO', 'AD BLUE'
    }

    for archivo in archivos:
        path = os.path.join(DIR_CONSOLIDADOS, archivo)
        try:
            df = pd.read_excel(path, sheet_name='B.D', engine='openpyxl')
            df.columns = df.columns.str.strip().str.upper()
            columnas_actuales = set(df.columns)

            if columnas_requeridas.issubset(columnas_actuales):
                dataframes.append(df)
            else:
                faltantes = sorted(columnas_requeridas - columnas_actuales)
                print(f"[ADVERTENCIA] Archivo omitido '{archivo}': faltan columnas {faltantes}")

        except Exception as e:
            print(f"[ADVERTENCIA] Error leyendo '{archivo}': {e}")

    if dataframes:
        print(f"[OK] Se cargaron {len(dataframes)} archivos consolidados correctamente.")
        return pd.concat(dataframes, ignore_index=True)
    else:
        print("[ERROR] No se encontraron archivos consolidados validos.")
        return None


def procesar_datos(df_consolidados, rangos_df):
    """Limpia, cruza datos, calcula acumulados y prepara la tabla exportable."""
    df = df_consolidados.copy()
    df.columns = df.columns.str.strip().str.upper()

    # Estandarizar nombres y tipos
    df.rename(columns={"ODOMETRO": "KILOMETROS"}, inplace=True)
    df['PATENTE'] = df['PATENTE'].astype(str).str.strip().str.upper()
    df['N INTERNO'] = df['NUMERO INTERNO'].astype(str).str.strip()
    df['KILOMETROS'] = pd.to_numeric(df['KILOMETROS'], errors='coerce')
    df['LITROS'] = pd.to_numeric(df['LITROS'], errors='coerce')

    # Formatear fechas
    df['FECHA REAL'] = pd.to_datetime(df['FECHA REAL'], errors='coerce').dt.strftime('%d-%m-%Y')
    df['FECHA PLANILLA'] = pd.to_datetime(df['FECHA PLANILLA'], errors='coerce').dt.strftime('%d-%m-%Y')

    # Limpieza de hora
    df['HORA'] = df['HORA'].apply(limpiar_hora)

    # Ordenar datos para calculo de KMACC
    df = df.sort_values(by=['PATENTE', 'FECHA REAL', 'HORA'], na_position='last')

    # Calculos iniciales de referencia en Python
    df['KMACC'] = df.groupby('PATENTE')['KILOMETROS'].diff().fillna(0)
    df['RENDIMIENTO'] = (df['KMACC'] / df['LITROS']).round(2)

    # Unir con tabla de Rangos
    rangos_df['PATENTE'] = rangos_df['PATENTE'].astype(str).str.strip().str.upper()
    rangos_df['N INTERNO'] = rangos_df['N INTERNO'].astype(str).str.strip()

    df_final = pd.merge(
        df,
        rangos_df[['PATENTE', 'N INTERNO', 'MODELO', 'NORMA', 'DESDE', 'HASTA']],
        how='left',
        on=['PATENTE', 'N INTERNO']
    )

    # Evaluacion de la columna REV (Revision)
    def evaluar_revision(row):
        if pd.notna(row['RENDIMIENTO']) and pd.notna(row['DESDE']) and row['RENDIMIENTO'] < row['DESDE']:
            return 'BR'
        elif pd.notna(row['RENDIMIENTO']) and pd.notna(row['HASTA']) and row['RENDIMIENTO'] > row['HASTA']:
            return 'CI'
        return 0

    df_final['REV'] = df_final.apply(evaluar_revision, axis=1)

    # Reordenar columnas finales
    columnas_finales = [
        'TERMINAL', 'FOLIO', 'FECHA PLANILLA', 'FECHA REAL', 'N INTERNO', 'PATENTE',
        'ROL', 'KILOMETROS', 'KMACC', 'LITROS', 'HORA', 'BOMBERO', 'TURNO',
        'RENDIMIENTO', 'MODELO', 'NORMA', 'DESDE', 'HASTA', 'REV', 'AD BLUE'
    ]
    
    return df_final[columnas_finales]


def exportar_excel_con_formato(df):
    """Crea el archivo final en Excel agregando formulas vivas y formato condicional."""
    wb = Workbook()
    wb.remove(wb.active)  # Eliminar hoja por defecto
    ws = wb.create_sheet("Reporte")

    # Estilos de relleno
    rojo = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    azul = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
    header_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")

    # Escribir filas y formato de encabezado
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        ws.append(row)
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if r_idx == 1:
                cell.font = Font(bold=True)
                cell.fill = header_fill

    # Mapeo de columnas a letras de Excel
    col_map = {cell.value: cell.column_letter for cell in ws[1]}
    max_row = ws.max_row

    # Inyeccion de formulas dinamicas por fila
    for row in range(2, max_row + 1):
        prev = row - 1

        # Formula Excel para KMACC
        ws[f"{col_map['KMACC']}{row}"] = (
            f'=IF(AND({col_map["PATENTE"]}{row}={col_map["PATENTE"]}{prev}, '
            f'ISNUMBER({col_map["KILOMETROS"]}{row}), ISNUMBER({col_map["KILOMETROS"]}{prev})), '
            f'{col_map["KILOMETROS"]}{row}-{col_map["KILOMETROS"]}{prev}, "")'
        )

        # Formula Excel para RENDIMIENTO
        ws[f"{col_map['RENDIMIENTO']}{row}"] = (
            f'=IFERROR(ROUND({col_map["KMACC"]}{row}/{col_map["LITROS"]}{row}, 2), "")'
        )

    # Formato condicional: KMACC es 0 o KILOMETROS vacio (Azul)
    ws.conditional_formatting.add(
        f"{col_map['KMACC']}2:{col_map['KMACC']}{max_row}",
        FormulaRule(
            formula=[f'OR({col_map["KMACC"]}2=0, ISBLANK({col_map["KILOMETROS"]}2))'],
            fill=azul
        )
    )

    # Formato condicional: RENDIMIENTO < DESDE (Rojo)
    ws.conditional_formatting.add(
        f"{col_map['RENDIMIENTO']}2:{col_map['RENDIMIENTO']}{max_row}",
        FormulaRule(
            formula=[
                f'AND(ISNUMBER({col_map["RENDIMIENTO"]}2),ISNUMBER({col_map["DESDE"]}2),{col_map["RENDIMIENTO"]}2<{col_map["DESDE"]}2)'
            ],
            fill=rojo
        )
    )

    # Formato condicional: RENDIMIENTO > HASTA (Amarillo)
    ws.conditional_formatting.add(
        f"{col_map['RENDIMIENTO']}2:{col_map['RENDIMIENTO']}{max_row}",
        FormulaRule(
            formula=[
                f'AND(ISNUMBER({col_map["RENDIMIENTO"]}2),ISNUMBER({col_map["HASTA"]}2),{col_map["RENDIMIENTO"]}2>{col_map["HASTA"]}2)'
            ],
            fill=amarillo
        )
    )

    wb.save(ARCHIVO_SALIDA)
    print(f"[OK] Archivo final generado con exito: {ARCHIVO_SALIDA}")


def main():
    """Flujo principal de ejecucion."""
    print("Iniciando proceso de consolidacion de rendimientos...\n")
    crear_carpetas()

    df_rangos = cargar_rangos()
    df_consolidados = cargar_consolidados()

    if df_rangos is None or df_consolidados is None or df_consolidados.empty:
        print("\n[ERROR] No se puede continuar. Revisa las carpetas de entrada.")
        return

    df_exportable = procesar_datos(df_consolidados, df_rangos)

    if df_exportable is not None:
        try:
            exportar_excel_con_formato(df_exportable)
            print("\nProceso completado exitosamente.")
        except Exception as e:
            print(f"\n[ERROR] Ocurrio un problema al aplicar formato en Excel: {e}")
    else:
        print("\n[ADVERTENCIA] No se procesaron datos para la exportacion.")


if __name__ == "__main__":
    main()